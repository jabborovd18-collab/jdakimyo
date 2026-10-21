# -*- coding: utf-8 -*-
"""
Mustaqil ta'lim hujjatlari uchun UMUMIY qurilma: bitta manbadan .docx va .pdf.

Nega alohida modul: har bir mavzu skripti o'zida nusxa ko'tarsa, uslub
(shrift, maydon, jadval rangi) va formulalarni qurish mantig'i vaqt o'tishi
bilan ajralib ketadi — bir joyda tuzatilgan xato boshqasida qoladi.

Formulalar bir marta OMML mini-DSL da yoziladi:
  * .docx — Word'ning tabiiy "Tenglama"si (OMML) sifatida kiritiladi;
  * .pdf  — o'sha OMML daraxti LaTeX ga aylantirilib, matplotlib mathtext
            orqali rasmga chiziladi va reportlab sahifasiga joylanadi.
Nega ikki backend: sandboxda Word/LibreOffice yo'q, docx→pdf konvertor yo'q;
lekin matnni ikki marta yozish mumkin emas (yagona manba qoidasi).
"""
import io
import os
import re
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from reportlab.lib.units import cm as _RCM

FONT = "Times New Roman"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


# ============================================================================
# OMML mini-DSL
# ============================================================================
def _el(tag):
    return etree.Element("{%s}%s" % (M, tag))


def _wrap(tag, items):
    el = _el(tag)
    if not isinstance(items, (list, tuple)):
        items = [items]
    for it in items:
        el.append(it)
    return el


def T(text, plain=False):
    """Matematik run. plain=True — tik shrift (ln, cos, birliklar, sonlar)."""
    r = _el("r")
    if plain:
        rpr = etree.SubElement(r, "{%s}rPr" % M)
        sty = etree.SubElement(rpr, "{%s}sty" % M)
        sty.set("{%s}val" % M, "p")
    t = etree.SubElement(r, "{%s}t" % M)
    t.set(XML_SPACE, "preserve")
    t.text = text
    return r


def P(text):
    """Tik (plain) matn: sonlar, 'ln', 'cos', qisqartma indekslar."""
    return T(text, plain=True)


def V(text):
    """O'zgaruvchi (kursiv — Word'ning standart matematik uslubi)."""
    return T(text)


def seq(*items):
    return list(items)


def frac(num, den):
    return _wrap("f", [_wrap("num", num), _wrap("den", den)])


def sub(base, s):
    return _wrap("sSub", [_wrap("e", base), _wrap("sub", s)])


def sup(base, s):
    return _wrap("sSup", [_wrap("e", base), _wrap("sup", s)])


def subsup(base, s, p_):
    return _wrap("sSubSup", [_wrap("e", base), _wrap("sub", s), _wrap("sup", p_)])


def par(items, left="(", right=")"):
    """Avtomatik o'lchamli qavslar."""
    d = _el("d")
    dpr = etree.SubElement(d, "{%s}dPr" % M)
    b = etree.SubElement(dpr, "{%s}begChr" % M)
    b.set("{%s}val" % M, left)
    e = etree.SubElement(dpr, "{%s}endChr" % M)
    e.set("{%s}val" % M, right)
    d.append(_wrap("e", items))
    return d


def sqrt(items):
    rad = _el("rad")
    rpr = etree.SubElement(rad, "{%s}radPr" % M)
    dh = etree.SubElement(rpr, "{%s}degHide" % M)
    dh.set("{%s}val" % M, "1")
    rad.append(_el("deg"))
    rad.append(_wrap("e", items))
    return rad


def bar(items):
    """Ustki chiziq (o'rtacha qiymat, v̄ kabi)."""
    acc = _el("acc")
    apr = etree.SubElement(acc, "{%s}accPr" % M)
    chr_ = etree.SubElement(apr, "{%s}chr" % M)
    chr_.set("{%s}val" % M, "\u0305")
    acc.append(_wrap("e", items))
    return acc


def omath(*items):
    om = _el("oMath")
    for it in items:
        if isinstance(it, (list, tuple)):
            for x in it:
                om.append(x)
        else:
            om.append(it)
    return om


# ============================================================================
# OMML → LaTeX (PDF uchun). Faqat DSL ishlatadigan konstruksiyalar qo'llanadi.
# ============================================================================
_SYM = {
    "α": r"\alpha", "β": r"\beta", "γ": r"\gamma", "δ": r"\delta", "ε": r"\varepsilon",
    "ζ": r"\zeta", "η": r"\eta", "θ": r"\theta", "κ": r"\kappa", "λ": r"\lambda",
    "μ": r"\mu", "ν": r"\nu", "ξ": r"\xi", "π": r"\pi", "ρ": r"\rho", "σ": r"\sigma",
    "τ": r"\tau", "φ": r"\varphi", "χ": r"\chi", "ψ": r"\psi", "ω": r"\omega",
    "Γ": r"\Gamma", "Δ": r"\Delta", "Σ": r"\Sigma", "Ω": r"\Omega",
    "∞": r"\infty", "≈": r"\approx", "∝": r"\propto", "·": r"\cdot", "−": "-",
    "→": r"\rightarrow", "ħ": r"\hbar", "°": r"\circ", "±": r"\pm", "≤": r"\leq",
    "≥": r"\geq", "√": r"\surd",
    # Nega: mathtext da "%" izoh belgisi, "⁺/⁻" unicode indekslarini esa chizmaydi.
    "%": r"\%", "⁺": "^{+}", "⁻": "^{-}",
}


def _tex_text(text, plain):
    out = []
    buf = ""

    def flush():
        nonlocal buf
        if buf:
            out.append(r"\mathrm{%s}" % buf if plain else buf)
            buf = ""

    for ch in text:
        if ch in _SYM:
            flush()
            out.append(_SYM[ch] + " ")
        elif ch == " ":
            flush()
            out.append(r"\ ")
        elif ch in "[]":
            flush()
            out.append(ch)
        elif ch == ",":
            flush()
            out.append("{,}")
        else:
            buf += ch
    flush()
    return "".join(out)


def _tex(el, mode="text"):
    """mode: 'display' — alohida qator (katta kasr), 'text' — jadval katagi
    (oddiy kasr), 'inline' — matn ichida (kasr qiya chiziq bilan, aks holda
    rasm qator balandligidan oshib keyingi qator ustiga tushadi)."""
    tag = etree.QName(el).localname
    kids = {etree.QName(c).localname: c for c in el}
    inner = lambda t: "".join(_tex(c, mode) for c in kids[t]) if t in kids else ""
    if tag == "r":
        plain = el.find("{%s}rPr/{%s}sty" % (M, M)) is not None
        t = el.find("{%s}t" % M)
        return _tex_text(t.text or "", plain)
    if tag == "f":
        if mode == "inline":
            return r"{}^{%s}\!/\!{}_{%s}" % (inner("num"), inner("den"))
        return r"\%s{%s}{%s}" % ("dfrac" if mode == "display" else "frac", inner("num"), inner("den"))
    if tag == "sSub":
        return "{%s}_{%s}" % (inner("e"), inner("sub"))
    if tag == "sSup":
        return "{%s}^{%s}" % (inner("e"), inner("sup"))
    if tag == "sSubSup":
        return "{%s}_{%s}^{%s}" % (inner("e"), inner("sub"), inner("sup"))
    if tag == "d":
        return r"\left(%s\right)" % inner("e")
    if tag == "rad":
        return r"\sqrt{%s}" % inner("e")
    if tag == "acc":
        return r"\bar{%s}" % inner("e")
    if tag == "oMath":
        return "".join(_tex(c, mode) for c in el)
    return ""


def omath_to_tex(om, mode="text"):
    return "$" + _tex(om, mode) + "$"


# ============================================================================
# PDF backend (reportlab + matplotlib mathtext)
# ============================================================================
class _Pdf:
    def __init__(self):
        import matplotlib
        matplotlib.use("Agg")
        from matplotlib import rcParams
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.fonts import addMapping

        rcParams["mathtext.fontset"] = "dejavuserif"  # matn shrifti bilan bir xil qalinlik
        self.colors = colors
        fdir = os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")
        # Nega DejaVu Serif: sandboxda Times yo'q; DejaVu lotin/kirill/yunon,
        # pastki indeks raqamlari va strelkalarni to'liq qoplaydi.
        for name, f in (("Serif", "DejaVuSerif.ttf"), ("Serif-Bold", "DejaVuSerif-Bold.ttf"),
                        ("Serif-Italic", "DejaVuSerif-Italic.ttf"),
                        ("Serif-BoldItalic", "DejaVuSerif-BoldItalic.ttf")):
            pdfmetrics.registerFont(TTFont(name, os.path.join(fdir, f)))
        addMapping("Serif", 0, 0, "Serif")
        addMapping("Serif", 1, 0, "Serif-Bold")
        addMapping("Serif", 0, 1, "Serif-Italic")
        addMapping("Serif", 1, 1, "Serif-BoldItalic")

        base = dict(fontName="Serif", fontSize=11, leading=15, spaceAfter=4)
        self.st = {
            "body": ParagraphStyle("body", alignment=TA_JUSTIFY, **base),
            "center": ParagraphStyle("center", alignment=TA_CENTER, **base),
            "right": ParagraphStyle("right", alignment=TA_RIGHT, **base),
            "h1": ParagraphStyle("h1", fontName="Serif-Bold", fontSize=13.5, leading=17,
                                 spaceBefore=10, spaceAfter=4, keepWithNext=1,
                                 textColor=colors.HexColor("#1F3A5F")),
            "h2": ParagraphStyle("h2", fontName="Serif-Bold", fontSize=12, leading=15,
                                 spaceBefore=8, spaceAfter=3, keepWithNext=1,
                                 textColor=colors.HexColor("#1F3A5F")),
            "bullet": ParagraphStyle("bullet", alignment=TA_JUSTIFY, leftIndent=18,
                                     bulletIndent=6, spaceAfter=2, **{k: v for k, v in base.items() if k != "spaceAfter"}),
            "cell": ParagraphStyle("cell", fontName="Serif", fontSize=9.5, leading=12),
            "cellh": ParagraphStyle("cellh", fontName="Serif-Bold", fontSize=9.5, leading=12,
                                    alignment=TA_CENTER),
        }
        self.flow = []
        self.tmp = tempfile.mkdtemp(prefix="omml-pdf-")
        self._n = 0

    # -- formulani rasmga chizish ----------------------------------------------
    def math_png(self, om, fontsize, mode="text"):
        import matplotlib.pyplot as plt
        from PIL import Image as PILImage

        tex = omath_to_tex(om, mode)
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, tex, fontsize=fontsize)
        self._n += 1
        path = os.path.join(self.tmp, "eq%04d.png" % self._n)
        dpi = 400
        fig.savefig(path, dpi=dpi, bbox_inches="tight", pad_inches=0.01, transparent=True)
        plt.close(fig)
        w, h = PILImage.open(path).size
        return path, w * 72.0 / dpi, h * 72.0 / dpi

    def _depth_pt(self, tex, fontsize):
        """Formulaning bazaviy chiziqdan pastga tushgan qismi (punkt)."""
        from matplotlib.font_manager import FontProperties
        from matplotlib.mathtext import MathTextParser
        parsed = MathTextParser("agg").parse(tex, dpi=72, prop=FontProperties(size=fontsize))
        return parsed.depth

    def inline_img(self, om, fontsize=11.5):
        path, w, h = self.math_png(om, fontsize, "inline")
        # Nega valign manfiy son: rasm pastini matn bazaviy chizig'idan formula
        # chuqurligiga tushiramiz — shunda harflar matn bilan bir chiziqda turadi,
        # kasrlar esa tabiiy ravishda ikki tomonga chiqadi.
        depth = self._depth_pt(omath_to_tex(om, "inline"), fontsize)
        return '<img src="%s" width="%.1f" height="%.1f" valign="%.1f"/>' % (path, w, h, -depth - 0.7)

    def block_img(self, om, fontsize=13, mode="display"):
        from reportlab.platypus import Image
        path, w, h = self.math_png(om, fontsize, mode)
        return Image(path, width=w, height=h)


def rcm(x):
    """reportlab uchun santimetr. Nega alohida: docx.shared.Cm EMU qaytaradi,
    reportlab esa punkt kutadi — aralashtirilsa sahifa ramkasi manfiy bo'ladi."""
    return x * _RCM


def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ============================================================================
# Hujjat yig'uvchi — har bir chaqiriq ikkala backendga yoziladi
# ============================================================================
class Hujjat:
    def __init__(self, pdf=True):
        self.doc = Document()
        self.eq_no = 0
        self.pdf = _Pdf() if pdf else None
        for s in self.doc.sections:
            s.top_margin = Cm(2)
            s.bottom_margin = Cm(2)
            s.left_margin = Cm(2.5)
            s.right_margin = Cm(1.5)

        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.15

        for name, size in (("Heading 1", 14), ("Heading 2", 12.5)):
            h = self.doc.styles[name]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)

    # -- ichki: rich qismlarni PDF markup satriga aylantirish -----------------
    def _markup(self, parts):
        out = []
        for item in parts:
            if isinstance(item, tuple):
                text, fmt = item
                s = _esc(text)
                if "b" in fmt:
                    s = "<b>%s</b>" % s
                if "i" in fmt:
                    s = "<i>%s</i>" % s
                out.append(s)
            else:
                out.append(self.pdf.inline_img(item) if self.pdf else "")
        return "".join(out)

    def _pdf_para(self, markup, style="body"):
        if self.pdf:
            from reportlab.platypus import Paragraph
            self.pdf.flow.append(Paragraph(markup, self.pdf.st[style]))

    # -- matn ----------------------------------------------------------------
    def h(self, text, level=1):
        self._pdf_para(_esc(text), "h1" if level == 1 else "h2")
        return self.doc.add_heading(text, level=level)

    def p(self, text, bold=False, italic=False, align="j", size=None, after=None):
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if align == "c":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "j":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if after is not None:
            para.paragraph_format.space_after = Pt(after)

        fmt = ("b" if bold else "") + ("i" if italic else "")
        m = self._markup([(text, fmt)])
        if size:
            m = '<font size="%s">%s</font>' % (size - 1, m)
        self._pdf_para(m, "center" if align == "c" else "body")
        return para

    def rich(self, parts, align="j", para=None, _pdf=True):
        """parts: (matn, 'b'|'i'|'') kortejlar yoki OMML element (inline)."""
        if para is None:
            para = self.doc.add_paragraph()
        for item in parts:
            if isinstance(item, tuple):
                text, fmt = item
                r = para.add_run(text)
                r.bold = "b" in fmt
                r.italic = "i" in fmt
            else:
                para._p.append(item)
        if align == "j":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if _pdf:
            self._pdf_para(self._markup(parts))
        return para

    def equation(self, om, where=None):
        """Raqamlangan, markazlashgan tenglama (chegarasiz 2 ustunli jadval).
        where — "bu yerda ..." izohi uchun rich() qismlari."""
        self.eq_no += 1
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t.rows[0].cells
        c0.width = Cm(14.5)
        c1.width = Cm(2.0)
        c0.paragraphs[0]._p.append(om)
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c1.paragraphs[0].add_run("(%d)" % self.eq_no)
        c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for c in (c0, c1):
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
            c.paragraphs[0].paragraph_format.space_before = Pt(2)

        if self.pdf:
            from reportlab.platypus import Paragraph, Table, TableStyle
            img = self.pdf.block_img(om)
            tbl = Table([[img, Paragraph("(%d)" % self.eq_no, self.pdf.st["right"])]],
                        colWidths=[rcm(14.5), rcm(2.0)])
            tbl.setStyle(TableStyle([
                ("ALIGN", (0, 0), (0, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            self.pdf.flow.append(tbl)

        if where:
            para = self.doc.add_paragraph()
            para.add_run("bu yerda ")
            self.rich(where, para=para, _pdf=False)
            para.paragraph_format.space_after = Pt(6)
            self._pdf_para(self._markup([("bu yerda ", "")] + list(where)))
        return self.eq_no

    def bullets(self, items, style="List Bullet"):
        numbered = style == "List Number"
        for n, it in enumerate(items, 1):
            # Nega qo'lda raqam: python-docx "List Number" butun hujjat bo'ylab
            # raqamni davom ettiradi (Reja 5 dan boshlanib ketadi).
            para = self.doc.add_paragraph(style="List Bullet" if not numbered else None)
            if numbered:
                para.paragraph_format.left_indent = Cm(0.75)
                para.paragraph_format.first_line_indent = Cm(-0.5)
                para.add_run("%d. " % n)
            if isinstance(it, tuple):
                parts = [(it[0], "b")] + (it[1] if isinstance(it[1], list) else [(it[1], "")])
            elif isinstance(it, list):
                parts = it
            else:
                parts = [(it, "")]
            self.rich(parts, para=para, _pdf=False)
            para.paragraph_format.space_after = Pt(2)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            if self.pdf:
                from reportlab.platypus import Paragraph
                self.pdf.flow.append(Paragraph(self._markup(parts), self.pdf.st["bullet"],
                                               bulletText=("%d." % n) if numbered else "•"))

    @staticmethod
    def _shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def table(self, headers, rows, widths=None):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, hd in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(hd)
            r.bold = True
            r.font.size = Pt(11)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._shade(c, "DCE6F1")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                if isinstance(v, str):
                    r = para.add_run(v)
                    r.font.size = Pt(11)
                else:  # OMML
                    para._p.append(v)
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

        if self.pdf:
            from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
            data = [[Paragraph(_esc(hd), self.pdf.st["cellh"]) for hd in headers]]
            for row in rows:
                data.append([Paragraph(_esc(v), self.pdf.st["cell"]) if isinstance(v, str)
                             else self.pdf.block_img(v, fontsize=11, mode="text") for v in row])
            tbl = Table(data, colWidths=[rcm(w) for w in widths] if widths else None, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, self.pdf.colors.black),
                ("ALIGN", (0, 1), (-1, -1), "CENTER"),
                ("BACKGROUND", (0, 0), (-1, 0), self.pdf.colors.HexColor("#DCE6F1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            self.pdf.flow.append(tbl)
            self.pdf.flow.append(Spacer(1, 6))
        return t

    def title(self, fan, mavzu):
        self.p("MUSTAQIL TA'LIM TOPSHIRIG'I", bold=True, align="c", size=13, after=2)
        self.p("Fan: " + fan, italic=True, align="c", size=11, after=10)
        self.p(mavzu, bold=True, align="c", size=14, after=10)

    def refs(self, items):
        self.h("Foydalanilgan adabiyotlar", 1)
        for i, r in enumerate(items, 1):
            para = self.doc.add_paragraph()
            para.add_run(f"{i}. {r}").font.size = Pt(11)
            para.paragraph_format.space_after = Pt(2)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self._pdf_para('<font size="10">%s</font>' % _esc(f"{i}. {r}"))

    def save(self, path):
        self.doc.save(path)
        print("Saqlandi:", path, "| tenglamalar soni:", self.eq_no)
        if self.pdf:
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate
            pdf_path = os.path.splitext(path)[0] + ".pdf"

            def footer(canvas, doc_):
                canvas.saveState()
                canvas.setFont("Serif", 9)
                canvas.drawCentredString(A4[0] / 2, rcm(1.1), str(doc_.page))
                canvas.restoreState()

            sd = SimpleDocTemplate(pdf_path, pagesize=A4, leftMargin=rcm(2.5), rightMargin=rcm(1.5),
                                   topMargin=rcm(2), bottomMargin=rcm(2),
                                   title=os.path.basename(path), author="Mustaqil ta'lim")
            sd.build(self.pdf.flow, onFirstPage=footer, onLaterPages=footer)
            print("Saqlandi:", pdf_path)
